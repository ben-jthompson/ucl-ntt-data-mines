import requests
import mimetypes
import os
import fitz
import docx
import json
import pandas as pd
import geopandas as gpd
from langchain.schema import Document

# find file type for scraped documents
def detect_file_type(url: str) -> str:
    # file type address stripping
    file_type = os.path.splitext(os.path.basename(url.split("?")[0]))[-1].lower()
    if file_type in {'.pdf', '.docx', '.csv', '.shp', '.json', '.geojson'}:
        return file_type
    
    # fallback for different address
    try:
        response = requests.head(url, allow_redirects=True, timeout=5)
        content_type = response.headers.get("Content-Type", "").lower()
    except Exception as e:
        print(f"Could not get Content-Type for {url}: {e}")
        return None

    if "pdf" in content_type:
        return '.pdf'
    elif "csv" in content_type:
        return '.csv'
    elif "msword" in content_type or "wordprocessingml" in content_type:
        return '.docx'
    elif "json" in content_type and "geo" not in content_type:
        return '.json'
    elif "geojson" in content_type:
        return '.geojson'
    return None  



def download_file(url: str, save_dir="server/downloads"):
    headers = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/114.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.5"
}
    os.makedirs(save_dir, exist_ok=True)
    filename = os.path.basename(url.split("?")[0])  
    filepath = os.path.join(save_dir, filename)

    response = requests.get(url, headers=headers)
    if response.status_code == 200:
        with open(filepath, 'wb') as f:
            f.write(response.content)
        print(f"Downloaded: {filename}")
        return filepath
    else:
        print(f"Failed to download: {url}")
        print(f"[INFO] Status code: {response.status_code}")
        print(f"[INFO] Response headers: {response.headers}")
        return None
    
def document_to_dict(doc: Document) -> dict:
    return {
        "page_content": doc.page_content,
        "metadata": doc.metadata
    }

def dict_to_document(dict: dict) -> Document:
    return Document(page_content=dict["page_content"], metadata=dict["metadata"])

def extract_from_pdf(path: str, save_dir="server/downloads") -> str:
    filename = os.path.basename(path.split("?")[0])  
    filepath = os.path.join(save_dir, filename)
    file = fitz.open(filepath)
    return "\n\n".join(page.get_text() for page in file)

def extract_from_docx(filepath: str) -> str:
    file = docx.Document(filepath)
    paras = [para.text for para in file.paragraphs]
    return '\n\n'.join(paras)

def extract_from_csv(filepath: str) -> str:
    file = pd.read_csv(filepath)
    return file.to_string(index=False)


def extract_from_json(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)
    return json.dumps(data, indent=2)

def extract_from_geojson(filepath: str) -> str:
    gdf = gpd.read_file(filepath)
    return gdf.to_json(indent=2)

def extract_from_shp(filepath: str) -> str:
    gdf = gpd.read_file(filepath)
    return gdf.to_string(index=False)

def extract_from_source(filepath: str, doc: Document, save_dir='server/downloads') -> Document:
    parsing_method = {'.pdf': extract_from_pdf, 
                      '.docx': extract_from_docx, 
                      '.csv': extract_from_csv, 
                      '.shp': extract_from_shp,
                      '.json': extract_from_json, 
                      '.geojson': extract_from_geojson }
    # if document has been found during scraping:
    if save_dir == 'server/downloads':
        filename = os.path.basename(filepath.split("?")[0])  
        filepath = os.path.join(save_dir, filename)
    # else document has been uploaded by user, so filepath is okay

    # iterate through keys to find appropriate extraction method
    for key, func in parsing_method.items():
        if filepath.endswith(key):
            # edge case: if doc
            doc.page_content=func(filepath)
            print('Page content for ', key, ' is ', doc.page_content[:100])
            return doc

